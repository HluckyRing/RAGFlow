# -*- coding: utf-8 -*-
"""文件加载器回归：CSV 多编码、TXT、PPTX 内容抽取，以及不支持格式的报错。"""
import io

import pytest
from pptx import Presentation
from pptx.util import Inches

import src.loaders as loaders


class FakeUpload:
    """模拟 FastAPI 的 UploadFile（构造后交给 _unwrap 处理）。

    除了 seek/read 还要给 tell/seekable/readable：python-pptx 读 pptx 时底层走
    zipfile，只实现 seek/read 的假对象会在 tell() 上炸掉（真实的
    UploadFile.file 是 SpooledTemporaryFile，这些方法都有）。
    """

    def __init__(self, data: bytes, name: str):
        self.name = name
        self._buf = io.BytesIO(data)

    def seek(self, *args):
        return self._buf.seek(*args)

    def tell(self):
        return self._buf.tell()

    def read(self, *args):
        return self._buf.read(*args)

    def seekable(self):
        return True

    def readable(self):
        return True


CSV_CASES = [
    # 简体样本
    ("utf-8", "姓名,分数\n张三,90\n", "姓名 | 分数"),
    ("utf-8-sig", "姓名,分数\n张三,90\n", "姓名 | 分数"),
    ("gbk", "姓名,分数\n张三,90\n", "姓名 | 分数"),
    ("gb2312", "姓名,分数\n张三,90\n", "姓名 | 分数"),
    # 繁体样本：Big5 编不出简体「数」字，所以单独准备一份
    ("big5", "姓名,分數\n張三,90\n", "姓名 | 分數"),
]


@pytest.mark.parametrize("encoding,text,first_line", CSV_CASES)
def test_load_csv_supports_common_encodings(encoding, text, first_line):
    out = loaders.load_file(FakeUpload(text.encode(encoding), "成绩.csv"))
    assert "90" in out
    # BOM 不能粘在第一个单元格上
    assert out.splitlines()[0] == first_line


def test_load_txt_still_supports_gbk():
    out = loaders.load_file(FakeUpload("中文内容\n".encode("gbk"), "a.txt"))
    assert "中文内容" in out


def test_unsupported_extension_raises():
    with pytest.raises(ValueError):
        loaders.load_file(FakeUpload(b"x", "a.bin"))


# ── PPTX 支持（2026-10-03 新增）──
# 样本用 python-pptx 在内存里现造，不往仓库塞二进制文件。


def _build_pptx():
    """造一份 2 页演示：第 1 页标题+正文，第 2 页表格+备注+组合形状。"""
    prs = Presentation()

    s1 = prs.slides.add_slide(prs.slide_layouts[5])            # Title Only
    s1.shapes.title.text = "年度总结"
    box = s1.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
    tf = box.text_frame
    tf.text = "第一点：营收增长"
    tf.add_paragraph().text = "第二点：成本下降"

    s2 = prs.slides.add_slide(prs.slide_layouts[6])            # Blank
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "ROE"
    table.cell(1, 1).text = "12%"
    # 备注常是真正的讲稿，问答命中率比页面上那几行字高
    s2.notes_slide.notes_text_frame.text = "这里重点讲成本控制"
    group = s2.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)).text_frame.text = "组合里的文字"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_load_pptx_extracts_page_text_with_page_headers():
    out = loaders.load_file(FakeUpload(_build_pptx(), "年度汇报.pptx"))
    assert "=== 第 1 页 ===" in out
    assert "=== 第 2 页 ===" in out
    assert "年度总结" in out
    assert "第一点：营收增长" in out
    assert "第二点：成本下降" in out


def test_load_pptx_extracts_tables_with_pipe_separator():
    """表格沿用 docx/xlsx 的 " | " 拼接约定。"""
    out = loaders.load_file(FakeUpload(_build_pptx(), "年度汇报.pptx"))
    assert "指标 | 数值" in out
    assert "ROE | 12%" in out


def test_load_pptx_extracts_speaker_notes():
    out = loaders.load_file(FakeUpload(_build_pptx(), "年度汇报.pptx"))
    assert "[备注] 这里重点讲成本控制" in out


def test_load_pptx_descends_into_group_shapes():
    """组合形状里的文字不下钻就会整块丢掉。"""
    out = loaders.load_file(FakeUpload(_build_pptx(), "年度汇报.pptx"))
    assert "组合里的文字" in out


def test_load_pptx_pptm_uses_same_loader():
    assert loaders.LOADERS[".pptm"] is loaders.LOADERS[".pptx"]


def test_load_pptx_without_any_text_returns_empty():
    """整页没文字时不留页头，好让上层照旧判「文件内容为空」。"""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])                 # 一张空白页
    buf = io.BytesIO()
    prs.save(buf)
    out = loaders.load_file(FakeUpload(buf.getvalue(), "纯图片.pptx"))
    assert out.strip() == ""


def test_missing_python_pptx_gives_chinese_hint(monkeypatch):
    monkeypatch.setattr(loaders, "HAS_PPTX", False)
    with pytest.raises(ImportError, match="python-pptx"):
        loaders.load_file(FakeUpload(b"x", "a.pptx"))


def test_legacy_ppt_gets_actionable_message():
    """旧版二进制 .ppt 读不了，提示要可执行（另存为 .pptx），而不是笼统的「不支持」。"""
    with pytest.raises(ValueError, match="另存为"):
        loaders.load_file(FakeUpload(b"x", "旧稿.ppt"))
